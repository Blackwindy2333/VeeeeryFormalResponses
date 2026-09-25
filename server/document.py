"""公文渲染数据模型（服务端）：标题解析、纯文本导出、DOCX 辅助。"""

from __future__ import annotations

import re
from datetime import datetime
from typing import Any

from server.prompt import split_meta_and_body

META_JSON = re.compile(r"```(?:json)?\s*\n(\{[\s\S]*?\})\s*```", re.IGNORECASE)


def document_date() -> str:
    now = datetime.now()
    return f"{now.year}年{now.month}月{now.day}日"


def build_document_view(
    raw_content: str,
    config: dict[str, Any],
    doc_type: str = "报告",
    *,
    streaming: bool = False,
    usage: dict[str, Any] | None = None,
    elapsed_ms: int | None = None,
) -> dict[str, Any]:
    """构造前端/导出共用的公文视图模型。"""
    title, body = split_meta_and_body(raw_content)
    doc = config.get("document") or {}
    return {
        "title": title,
        "title_display": title or "（未识别标题）",
        "body": body,
        "masthead": doc.get("masthead") or "演示单位（示例）",
        "greeting": doc.get("greeting") or "尊敬的阅办人：",
        "closing": doc.get("closing") or "此致",
        "signature": doc.get("signature") or "智能助手",
        "role_title": doc.get("role_title") or "",
        "date": document_date(),
        "doc_type": doc_type,
        "streaming": streaming,
        "indent_paragraphs": doc.get("indent_paragraphs", True),
        "hierarchy_fonts": doc.get("hierarchy_fonts", True),
        "usage": usage,
        "elapsed_ms": elapsed_ms,
        "disclaimer": "说明：本文为演示界面生成内容，不具有法定公文效力。",
    }


def to_plain_text(view: dict[str, Any]) -> str:
    body = re.sub(r"```[\s\S]*?```", lambda m: m.group(0).split("\n", 1)[-1].rsplit("```", 1)[0], view.get("body") or "")
    body = re.sub(r"^#{1,6}\s+", "", body, flags=re.MULTILINE)
    body = re.sub(r"\*\*([^*]+)\*\*", r"\1", body)
    body = re.sub(r"`([^`]+)`", r"\1", body)
    lines = [
        view.get("masthead") or "",
        "————————————",
        view.get("title_display") or "",
        "",
        view.get("greeting") or "",
        "",
        body.strip(),
        "",
        view.get("closing") or "",
    ]
    if view.get("role_title"):
        lines.append(view["role_title"])
    lines.append(view.get("signature") or "")
    lines.append(view.get("date") or "")
    lines.append("")
    lines.append(f"（文种：{view.get('doc_type')}；演示生成，无法定效力）")
    return "\n".join(lines)


def to_docx_paragraphs(view: dict[str, Any]) -> list[tuple[str, str]]:
    """返回 (style_name, text) 列表，供 python-docx 写入。"""
    body = view.get("body") or ""
    # 粗略按空行分段
    paras: list[tuple[str, str]] = []
    paras.append(("Masthead", view.get("masthead") or ""))
    paras.append(("Title", view.get("title_display") or ""))
    paras.append(("Greeting", view.get("greeting") or ""))
    for block in re.split(r"\n\s*\n", body):
        text = block.strip()
        if not text:
            continue
        if text.startswith("#"):
            text = re.sub(r"^#{1,6}\s+", "", text)
            paras.append(("Heading", text.replace("\n", " ")))
        else:
            paras.append(("Body", text.replace("\n", "")))
    paras.append(("Closing", view.get("closing") or ""))
    if view.get("role_title"):
        paras.append(("Sign", view["role_title"]))
    paras.append(("Sign", view.get("signature") or ""))
    paras.append(("Sign", view.get("date") or ""))
    return paras


def try_export_docx(view: dict[str, Any], path: str) -> bool:
    """若环境有 python-docx 则导出，成功返回 True。"""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
    except ImportError:
        return False

    doc = Document()
    # 页边距接近公文
    for section in doc.sections:
        section.top_margin = _cm(3.7)
        section.bottom_margin = _cm(3.5)
        section.left_margin = _cm(2.8)
        section.right_margin = _cm(2.6)

    for style_name, text in to_docx_paragraphs(view):
        p = doc.add_paragraph()
        run = p.add_run(text)
        if style_name == "Masthead":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x8B, 0x1E, 0x1E)
            run.font.name = "SimSun"
        elif style_name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.bold = True
            run.font.size = Pt(16)
            run.font.name = "SimHei"
        elif style_name in {"Greeting", "Body", "Closing", "Sign"}:
            run.font.size = Pt(14)
            run.font.name = "FangSong"
            if style_name == "Sign":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif style_name == "Body":
                p.paragraph_format.first_line_indent = Pt(28)
        elif style_name == "Heading":
            run.bold = True
            run.font.size = Pt(14)
            run.font.name = "SimHei"
    doc.save(path)
    return True


def _cm(value: float):
    from docx.shared import Cm

    return Cm(value)
